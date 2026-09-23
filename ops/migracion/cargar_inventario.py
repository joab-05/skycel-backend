#!/usr/bin/env python3
"""
Carga el inventario existente del sistema anterior (skyceldb) en skycel-backend usando su API.

Por qué por la API y no con INSERT directo: así pasa por las mismas validaciones, la auditoría (Envers),
el historial de movimientos de inventario (ALTA / "Stock inicial") y la restricción única (codpro, codti).

Entradas (fuera de Git, contienen datos del negocio):
  --plan     plan_migracion.json  (lo genera preparar_plan.py)
  --nombres  nombres_articulos.xlsx  (opcional: nombres de artículo ya revisados por el usuario)

Uso típico:
  set SKYCEL_PASSWORD=...            (o --password-file ruta; NUNCA como argumento)
  python cargar_inventario.py --url https://skycelsys.ddns.net --usuario root --plan plan.json --nombres nombres.xlsx --solo-catalogos
  python cargar_inventario.py ... --tienda 8 --limite 50      (prueba corta)
  python cargar_inventario.py ...                              (carga completa; se puede interrumpir y reanudar)

Reglas de migración (acordadas con el usuario, 2026-09-23):
  * Accesorios: se CONSERVA el código viejo como codpro (es la etiqueta pegada al producto).
  * Equipos (celular/tablet/módem): entran por IMEI o número de serie, una unidad cada uno; el costo o precio
    distinto al del modelo queda como costo/precio propio de esa unidad.
  * Cada tienda conserva su propio costo, precio y proveedor.
  * El límite del servidor es de 100 peticiones por minuto por IP: el script se auto-regula (~90/min).
"""
import argparse, collections, getpass, json, os, sys, time, urllib.error, urllib.request

MAPA_TIENDAS = {1: 'Matriz Central', 2: 'Skycel Zocalo', 6: 'Skycel Mina', 7: 'Skycel Superche',
                8: 'Skycel Corpo', 9: 'Skycel Morelos', 10: 'Taller Superche'}
TIPO_CATEGORIA = {'Accesorio': 'ACCESORIO', 'Equipo (Celular)': 'CELULAR', 'Equipo (Tablet)': 'TABLET'}
TIPOS_EQUIPO = {'CELULAR', 'TABLET', 'MODEM'}


class Api:
    def __init__(self, base, intervalo=0.66):
        self.base = base.rstrip('/'); self.intervalo = intervalo; self.token = None; self._ultimo = 0.0

    def llamar(self, metodo, ruta, cuerpo=None, reintentos=4):
        for intento in range(reintentos):
            espera = self.intervalo - (time.time() - self._ultimo)
            if espera > 0: time.sleep(espera)
            self._ultimo = time.time()
            datos = json.dumps(cuerpo).encode('utf-8') if cuerpo is not None else None
            req = urllib.request.Request(self.base + ruta, data=datos, method=metodo)
            req.add_header('Content-Type', 'application/json; charset=utf-8')
            if self.token: req.add_header('Authorization', 'Bearer ' + self.token)
            try:
                with urllib.request.urlopen(req, timeout=60) as r:
                    txt = r.read().decode('utf-8')
                    return r.status, (json.loads(txt) if txt else None)
            except urllib.error.HTTPError as e:
                txt = e.read().decode('utf-8', 'replace')
                if e.code == 429:
                    time.sleep(25); continue                     # límite de peticiones: esperar y reintentar
                if e.code >= 500 and intento < reintentos - 1:
                    time.sleep(3); continue
                try: return e.code, json.loads(txt)
                except Exception: return e.code, {'message': txt[:300]}
            except (urllib.error.URLError, TimeoutError) as e:
                if intento < reintentos - 1: time.sleep(5); continue
                return 0, {'message': f'sin conexión: {e}'}
        return 0, {'message': 'reintentos agotados'}


def desenvolver(j):
    """Las respuestas de catálogos vienen como {success, data}; las demás, directas."""
    return j['data'] if isinstance(j, dict) and 'data' in j and 'success' in j else j


def mensaje(j):
    if isinstance(j, dict):
        return str(j.get('description') or j.get('message') or j.get('title') or j.get('detail') or j)[:300]
    return str(j)[:300]


def preparar_catalogos(api, plan, colores_req, provs_req):
    """Crea (si faltan) tiendas, categorías, colores y proveedores. Devuelve los mapas de ids."""
    st, j = api.llamar('GET', '/api/tiendas')
    tiendas = {t['nombre'].strip().lower(): t['codti'] for t in (j or [])} if st == 200 else {}
    id_tienda = {}
    for viejo, nombre in MAPA_TIENDAS.items():
        cod = tiendas.get(nombre.lower())
        if cod is None:
            st, j = api.llamar('POST', '/api/tiendas', {'nombre': nombre, 'esAlmacen': False})
            if st not in (200, 201): sys.exit(f'No se pudo crear la tienda {nombre}: {mensaje(j)}')
            cod = j['codti']; print(f'  + tienda creada: {nombre} (codti {cod})')
        id_tienda[viejo] = cod

    st, j = api.llamar('GET', '/api/categorias')
    existentes = {(c['nombre'].strip().lower()): c for c in (j or [])} if st == 200 else {}
    for nombre, codigo, tipo_nuevo in plan['categorias']:
        tipo = TIPO_CATEGORIA[tipo_nuevo]
        c = existentes.get(nombre.lower())
        if c:
            if c.get('tipo') != tipo: sys.exit(f'La categoría "{nombre}" ya existe con tipo {c.get("tipo")}, se esperaba {tipo}.')
            continue
        st, j = api.llamar('POST', '/api/categorias', {'nombreCat': nombre, 'tipo': tipo, 'codigo': codigo})
        if st not in (200, 201): sys.exit(f'No se pudo crear la categoría {nombre}: {mensaje(j)}')
        print(f'  + categoría creada: {nombre} ({tipo}, {codigo})')

    st, j = api.llamar('GET', '/api/catalogos/colores')
    id_color = {c['nombre'].strip().lower(): c['id'] for c in (desenvolver(j) or [])} if st == 200 else {}
    for nombre in sorted(colores_req):
        if nombre.lower() in id_color: continue
        st, j = api.llamar('POST', '/api/catalogos/colores', {'nombre': nombre})
        if st not in (200, 201): sys.exit(f'No se pudo crear el color {nombre}: {mensaje(j)}')
        id_color[nombre.lower()] = desenvolver(j)['id']; print(f'  + color creado: {nombre}')

    st, j = api.llamar('GET', '/api/catalogos/proveedores')
    id_prov = {p['nombreCorto'].strip().lower(): p['id'] for p in (desenvolver(j) or [])} if st == 200 else {}
    for nombre in sorted(provs_req):
        if nombre.lower() in id_prov: continue
        st, j = api.llamar('POST', '/api/catalogos/proveedores', {'nombreCorto': nombre[:50]})
        if st not in (200, 201): sys.exit(f'No se pudo crear el proveedor {nombre}: {mensaje(j)}')
        id_prov[nombre.lower()] = desenvolver(j)['id']; print(f'  + proveedor creado: {nombre}')
    return id_tienda, id_color, id_prov


def moda(valores):
    return collections.Counter(valores).most_common(1)[0][0]


def titulo_equipo(s):
    return ' '.join(w.upper() if w.upper() in ('HTC', 'ZTE', 'TCL') else w.title() for w in s.split())


def construir_cargas(filas, id_tienda, id_color, id_prov):
    """Una carga (POST /api/productos) por accesorio y por grupo de equipos (tienda + artículo + color)."""
    cargas = []; equipos = collections.defaultdict(list)
    for r in filas:
        cod = id_tienda[r['codti']]
        idc = id_color.get(r['color'].lower()) if r['color'] else None
        idp = id_prov.get(r['proveedor'].lower()) if r['proveedor'] else None
        if r['tipo_viejo'] in TIPOS_EQUIPO:
            equipos[(cod, r['nombre'], r['color'] or '', r['categoria'])].append((r, idc, idp))
            continue
        cargas.append(dict(clave=f"{r['codigo']}|{cod}", cuerpo={
            'codpro': r['codigo'], 'codti': cod, 'stock': r['existencia'],
            'precioCompra': round(r['costo'], 2), 'precioVenta': round(r['publico'], 2),
            'idColor': idc, 'idProveedor': idp, 'stockMinimo': 0,
            'tipoMaster': 'ACCESORIO', 'categoriaMaster': r['categoria'],
            'nombreMaster': r['nombre'], 'descripcion': r['nombre'], 'compatibilidad': r['compat'] or None}))
    for (cod, nombre, color, cat), lst in equipos.items():
        r0 = lst[0][0]
        costo = moda([round(x[0]['costo'], 2) for x in lst]); precio = moda([round(x[0]['publico'], 2) for x in lst])
        unidades = []
        for r, _, _ in lst:
            u = {'imei': r['codigo']}
            if round(r['costo'], 2) != costo: u['costoUnitario'] = round(r['costo'], 2)
            if round(r['publico'], 2) != precio: u['precioVenta'] = round(r['publico'], 2)
            unidades.append(u)
        cargas.append(dict(clave=f"EQ|{nombre}|{color}|{cod}", cuerpo={
            'codti': cod, 'stock': len(lst), 'precioCompra': costo, 'precioVenta': precio,
            'idColor': lst[0][1], 'idProveedor': moda([x[2] for x in lst]), 'stockMinimo': 0,
            'tipoMaster': 'TABLET' if r0['tipo_viejo'] == 'TABLET' else 'CELULAR', 'categoriaMaster': cat,
            'nombreMaster': nombre, 'marca': titulo_equipo(r0['marca']), 'modelo': titulo_equipo(r0['modelo']),
            'unidades': unidades}))
    return cargas


def aplicar_nombres(plan, ruta_xlsx):
    """Sustituye el nombre propuesto por el que el usuario corrigió en la hoja Nombres (col. B; original en col. H)."""
    from openpyxl import load_workbook
    ws = load_workbook(ruta_xlsx, read_only=True)['Nombres']
    cambios = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[7] and r[1] and str(r[1]).strip() != str(r[7]).strip(): cambios[str(r[7]).strip()] = str(r[1]).strip()
    for x in plan['plan']:
        if x['nombre'] in cambios: x['nombre'] = cambios[x['nombre']]
    return len(cambios)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--url', required=True); ap.add_argument('--usuario', required=True)
    ap.add_argument('--password-file', help='archivo con la contraseña en la primera línea')
    ap.add_argument('--plan', required=True); ap.add_argument('--nombres')
    ap.add_argument('--tienda', type=int, help='solo esta tienda (código VIEJO)'); ap.add_argument('--limite', type=int)
    ap.add_argument('--solo-catalogos', action='store_true'); ap.add_argument('--progreso', default='progreso_carga.json')
    ap.add_argument('--errores', default='errores_carga.jsonl'); ap.add_argument('--intervalo', type=float, default=0.66)
    a = ap.parse_args()

    pw = os.environ.get('SKYCEL_PASSWORD')
    if a.password_file: pw = open(a.password_file, encoding='utf-8').readline().strip()
    if not pw: pw = getpass.getpass('Contraseña: ')
    plan = json.load(open(a.plan, encoding='utf-8'))
    if a.nombres: print(f'Nombres corregidos aplicados: {aplicar_nombres(plan, a.nombres)}')
    filas = [x for x in plan['plan'] if (a.tienda is None or x['codti'] == a.tienda) and x['codti'] in MAPA_TIENDAS]
    omitidas = [x for x in plan['plan'] if x['codti'] not in MAPA_TIENDAS and (a.tienda is None or x['codti'] == a.tienda)]
    if omitidas: print(f'AVISO: {len(omitidas)} filas de tiendas no migradas (cerradas) se omiten.')

    api = Api(a.url, a.intervalo)
    st, j = api.llamar('POST', '/api/auth/login', {'username': a.usuario, 'password': pw})
    if st != 200: sys.exit(f'Login fallido ({st}): {mensaje(j)}')
    api.token = j['token']; print(f'Sesión iniciada como {j.get("username")} ({j.get("rol")}).')

    colores = {x['color'] for x in plan['plan'] if x['color']}; provs = {x['proveedor'] for x in plan['plan'] if x['proveedor']}
    print('Preparando catálogos…')
    id_tienda, id_color, id_prov = preparar_catalogos(api, plan, colores, provs)
    if a.solo_catalogos: print('Catálogos listos.'); return

    cargas = construir_cargas(filas, id_tienda, id_color, id_prov)
    if a.limite: cargas = cargas[:a.limite]
    hechas = set(json.load(open(a.progreso)) if os.path.exists(a.progreso) else [])
    pend = [c for c in cargas if c['clave'] not in hechas]
    print(f'{len(cargas)} cargas ({len(cargas) - len(pend)} ya hechas). Tiempo estimado: {len(pend) * a.intervalo / 60:.0f} min.')
    ok = err = 0; t0 = time.time()
    for i, c in enumerate(pend, 1):
        st, j = api.llamar('POST', '/api/productos', c['cuerpo'])
        if st in (200, 201) or (st == 409 and 'Ya existe' in mensaje(j)):
            hechas.add(c['clave']); ok += 1
        else:
            err += 1
            with open(a.errores, 'a', encoding='utf-8') as f:
                f.write(json.dumps({'clave': c['clave'], 'status': st, 'error': mensaje(j), 'cuerpo': c['cuerpo']}, ensure_ascii=False) + '\n')
        if i % 50 == 0 or i == len(pend):
            json.dump(sorted(hechas), open(a.progreso, 'w'))
            print(f'  {i}/{len(pend)}  ok={ok} errores={err}  ({(time.time() - t0) / 60:.1f} min)', flush=True)
    print(f'Terminado: {ok} cargadas, {err} con error' + (f' (ver {a.errores})' if err else '') + '.')


if __name__ == '__main__':
    main()
